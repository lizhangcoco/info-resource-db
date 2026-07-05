#!/usr/bin/env python3
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from flask import Flask, render_template, jsonify, request
from datetime import datetime

from storage import database
from core.engine import get_engine
from core.analyzer import full_analysis, mark_recommendations
from core.trend import get_keyword_trends, batch_trends_to_echarts
from collectors import get_supported_platforms
from core.auth import generate_password_hash, verify_password, generate_jwt, decode_jwt
from storage.models import User, Supplier, SupplierProduct, RFQRecord
from storage.categories import CATEGORIES, get_main_categories, get_sub_categories


def create_app():
    app = Flask(__name__,
                template_folder=os.path.join(os.path.dirname(__file__), "web", "templates"),
                static_folder=os.path.join(os.path.dirname(__file__), "web", "static"))
    app.config["APPLICATION_ROOT"] = "/bijia"

    database.init_db()
    engine = get_engine()

    def get_current_user():
        token = request.headers.get("Authorization", "").replace("Bearer ", "")
        if not token:
            return None
        payload = decode_jwt(token)
        if not payload:
            return None
        return database.get_user_by_id(payload.get("user_id"))

    def require_auth(role=None):
        user = get_current_user()
        if not user:
            return jsonify({"error": "未登录"}), 401
        if user.status != "active":
            return jsonify({"error": "账号已停用"}), 403
        if role and user.role != role:
            return jsonify({"error": "权限不足"}), 403
        return None

    @app.route("/")
    def index():
        return render_template("index.html")

    @app.route("/login")
    def login_page():
        return render_template("login.html")

    @app.route("/admin")
    def admin_page():
        return render_template("admin.html")

    @app.route("/supplier-products")
    def supplier_products_page():
        return render_template("supplier_products.html")

    @app.route("/api/platforms")
    def api_platforms():
        return jsonify({
            "platforms": get_supported_platforms()
        })

    @app.route("/api/keywords")
    def api_keywords():
        keywords = database.get_keywords()
        return jsonify({"keywords": keywords})

    @app.route("/api/search", methods=["POST"])
    def api_search():
        auth_check = require_auth()
        if auth_check:
            return auth_check

        data = request.get_json() or {}
        keyword = data.get("keyword", "").strip()
        platforms = data.get("platforms") or get_supported_platforms()
        limit = int(data.get("limit", 15))

        if not keyword:
            return jsonify({"error": "关键词不能为空"}), 400

        record_id = engine.search_async(keyword, platforms, limit)
        return jsonify({"record_id": record_id, "status": "running"})

    @app.route("/api/search/<int:record_id>")
    def api_search_status(record_id):
        auth_check = require_auth()
        if auth_check:
            return auth_check

        status = engine.get_task_status(record_id)
        if status.get("status") == "completed" and "result" in status:
            result = status["result"]
            return jsonify(result)
        return jsonify(status)

    @app.route("/api/products")
    def api_products():
        auth_check = require_auth()
        if auth_check:
            return auth_check

        keyword = request.args.get("keyword", "").strip()
        platform = request.args.get("platform", "").strip() or None
        order_by = request.args.get("order_by", "price")
        sort = request.args.get("sort", "asc")
        limit = int(request.args.get("limit", 50))

        if not keyword:
            return jsonify({"products": [], "stats": {}})

        products = database.get_products_by_keyword(keyword, platform, order_by, sort, limit)
        products = mark_recommendations(products)
        stats = database.get_stats(keyword)

        supplier_products = database.get_supplier_products(keyword=keyword, limit=30)
        supplier_items = []
        for sp in supplier_products:
            supplier = database.get_supplier_by_id(sp.supplier_id)
            supplier_items.append({
                **sp.to_dict(),
                "supplier_name": supplier.name if supplier else "",
                "contact_phone": supplier.contact_phone if supplier else "",
                "platform": "supplier",
                "product_key": f"supplier_{sp.id}",
                "shop_name": supplier.name if supplier else "",
            })

        all_products = [p.to_dict() for p in products] + supplier_items
        if order_by == "price":
            all_products.sort(key=lambda x: x["price"], reverse=(sort.lower() == "desc"))

        return jsonify({
            "products": all_products,
            "stats": stats.to_dict(),
            "supplier_count": len(supplier_items),
        })

    @app.route("/api/products/<product_key>/trend")
    def api_product_trend(product_key):
        auth_check = require_auth()
        if auth_check:
            return auth_check

        days = int(request.args.get("days", 30))
        from core.trend import get_product_trend, trend_to_echarts
        trend = get_product_trend(product_key, days)
        return jsonify(trend_to_echarts(trend))

    @app.route("/api/trends")
    def api_trends():
        auth_check = require_auth()
        if auth_check:
            return auth_check

        keyword = request.args.get("keyword", "").strip()
        days = int(request.args.get("days", 30))
        limit = int(request.args.get("limit", 5))

        if not keyword:
            return jsonify({"trends": []})

        trends = get_keyword_trends(keyword, days, limit)
        return jsonify({"trends": batch_trends_to_echarts(trends)})

    @app.route("/api/stats")
    def api_stats():
        auth_check = require_auth()
        if auth_check:
            return auth_check

        keyword = request.args.get("keyword", "").strip()
        if not keyword:
            return jsonify({})
        stats = database.get_stats(keyword)
        return jsonify(stats.to_dict())

    @app.route("/api/history")
    def api_history():
        auth_check = require_auth()
        if auth_check:
            return auth_check

        keyword = request.args.get("keyword", "").strip() or None
        limit = int(request.args.get("limit", 20))
        records = database.get_search_records(keyword, limit)
        return jsonify({"records": [r.to_dict() for r in records]})

    @app.route("/api/upload", methods=["POST"])
    def api_upload():
        auth_check = require_auth()
        if auth_check:
            return auth_check

        try:
            data = request.get_json()
            if not data:
                return jsonify({"error": "请提供JSON数据"}), 400

            keyword = data.get("keyword", "").strip()
            products_data = data.get("products", [])

            if not keyword:
                return jsonify({"error": "关键词不能为空"}), 400

            if not products_data:
                return jsonify({"error": "商品数据不能为空"}), 400

            from storage.models import Product, PricePoint
            from core.cleaner import clean_products

            products = []
            now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            for i, p in enumerate(products_data):
                product = Product(
                    product_key=f"custom_{keyword}_{i}_{int(datetime.now().timestamp())}",
                    platform=p.get("platform", "custom"),
                    title=p.get("title", "未知商品"),
                    price=float(p.get("price", 0)),
                    url=p.get("url", "#"),
                    image_url=p.get("image_url", ""),
                    shop_name=p.get("shop_name", "自定义"),
                    shop_rating=float(p.get("shop_rating", 4.5)),
                    sales=int(p.get("sales", 0)),
                    keyword=keyword,
                )
                products.append(product)

            products = clean_products(products)
            database.batch_insert_products(products)

            price_points = []
            for p in products:
                price_points.append(PricePoint(
                    product_key=p.product_key,
                    price=p.price,
                    collected_at=now,
                    keyword=keyword,
                ))
            database.batch_insert_price_history(price_points)

            analysis = full_analysis(products)
            analysis["uploaded"] = True
            analysis["upload_count"] = len(products)

            return jsonify({
                "success": True,
                "keyword": keyword,
                "count": len(products),
                "products": analysis["products"],
                "stats": analysis["stats"],
            })

        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/api/upload/template")
    def api_upload_template():
        template = {
            "keyword": "商品关键词",
            "products": [
                {
                    "title": "商品名称",
                    "price": 1999.00,
                    "sales": 1000,
                    "shop_name": "店铺名",
                    "platform": "custom",
                    "url": "https://example.com/product"
                }
            ]
        }
        return jsonify(template)

    @app.route("/api/auth/login", methods=["POST"])
    def api_login():
        data = request.get_json() or {}
        username = data.get("username", "").strip()
        password = data.get("password", "").strip()

        if not username or not password:
            return jsonify({"error": "用户名或密码不能为空"}), 400

        user = database.get_user_by_username(username)
        if not user:
            return jsonify({"error": "用户名或密码错误"}), 401

        if not verify_password(password, user.password_hash):
            return jsonify({"error": "用户名或密码错误"}), 401

        if user.status != "active":
            return jsonify({"error": "账号已停用"}), 403

        token = generate_jwt(user.id, user.username, user.role)
        return jsonify({
            "success": True,
            "token": token,
            "user": user.to_dict(),
        })

    @app.route("/api/auth/register", methods=["POST"])
    def api_register():
        data = request.get_json() or {}
        username = data.get("username", "").strip()
        password = data.get("password", "").strip()
        email = data.get("email", "").strip()
        phone = data.get("phone", "").strip()
        company_name = data.get("company_name", "").strip()
        role = data.get("role", "buyer").strip()

        if not username or not password:
            return jsonify({"error": "用户名和密码不能为空"}), 400

        if len(password) < 6:
            return jsonify({"error": "密码长度至少6位"}), 400

        if role not in ("buyer", "supplier"):
            return jsonify({"error": "无效的角色类型"}), 400

        existing_user = database.get_user_by_username(username)
        if existing_user:
            return jsonify({"error": "用户名已存在"}), 400

        password_hash = generate_password_hash(password)
        supplier_id = None

        if role == "supplier":
            supplier_name = data.get("supplier_name", "").strip() or company_name
            if not supplier_name:
                return jsonify({"error": "供应商名称不能为空"}), 400

            supplier = Supplier(
                name=supplier_name,
                contact_name=data.get("contact_name", "").strip(),
                contact_phone=phone,
                contact_email=email,
                address=data.get("address", "").strip(),
                business_license=data.get("business_license", "").strip(),
                qualifications=data.get("qualifications", "").strip(),
                status="pending",
                remark="供应商自主注册，待审核",
            )
            supplier_id = database.create_supplier(supplier)

        user_status = "active" if role == "buyer" else "pending"
        user = User(
            username=username,
            password_hash=password_hash,
            email=email,
            phone=phone,
            role=role,
            company_name=company_name,
            supplier_id=supplier_id,
            status=user_status,
        )
        user_id = database.create_user(user)

        if role == "supplier":
            return jsonify({
                "success": True,
                "message": "注册成功，请等待管理员审核",
                "user": database.get_user_by_id(user_id).to_dict(),
            })

        token = generate_jwt(user_id, username, role)
        return jsonify({
            "success": True,
            "token": token,
            "user": database.get_user_by_id(user_id).to_dict(),
        })

    @app.route("/api/auth/me")
    def api_auth_me():
        user = get_current_user()
        if not user:
            return jsonify({"error": "未登录"}), 401
        return jsonify({"user": user.to_dict()})

    @app.route("/api/users", methods=["GET"])
    def api_users_list():
        auth_check = require_auth("admin")
        if auth_check:
            return auth_check

        role = request.args.get("role")
        status = request.args.get("status")
        limit = int(request.args.get("limit", 50))
        users = database.get_users(role, status, limit)
        return jsonify({"users": [u.to_dict() for u in users]})

    @app.route("/api/users/<int:user_id>", methods=["GET"])
    def api_users_get(user_id):
        auth_check = require_auth("admin")
        if auth_check:
            return auth_check

        user = database.get_user_by_id(user_id)
        if not user:
            return jsonify({"error": "用户不存在"}), 404
        return jsonify({"user": user.to_dict()})

    @app.route("/api/users/<int:user_id>", methods=["PUT"])
    def api_users_update(user_id):
        auth_check = require_auth("admin")
        if auth_check:
            return auth_check

        data = request.get_json() or {}
        update_fields = {}
        if "email" in data:
            update_fields["email"] = data["email"]
        if "phone" in data:
            update_fields["phone"] = data["phone"]
        if "company_name" in data:
            update_fields["company_name"] = data["company_name"]
        if "role" in data:
            update_fields["role"] = data["role"]
        if "status" in data:
            update_fields["status"] = data["status"]
        if "member_expire_at" in data:
            update_fields["member_expire_at"] = data["member_expire_at"]
        if "password" in data and data["password"]:
            update_fields["password_hash"] = generate_password_hash(data["password"])

        database.update_user(user_id, **update_fields)
        user = database.get_user_by_id(user_id)
        return jsonify({"success": True, "user": user.to_dict()})

    @app.route("/api/suppliers", methods=["GET"])
    def api_suppliers_list():
        auth_check = require_auth()
        if auth_check:
            return auth_check

        status = request.args.get("status")
        limit = int(request.args.get("limit", 50))
        suppliers = database.get_suppliers(status, limit)
        return jsonify({"suppliers": [s.to_dict() for s in suppliers]})

    @app.route("/api/suppliers", methods=["POST"])
    def api_suppliers_create():
        auth_check = require_auth("admin")
        if auth_check:
            return auth_check

        data = request.get_json() or {}
        supplier = Supplier(
            name=data.get("name", ""),
            contact_name=data.get("contact_name", ""),
            contact_phone=data.get("contact_phone", ""),
            contact_email=data.get("contact_email", ""),
            address=data.get("address", ""),
            business_license=data.get("business_license", ""),
            qualifications=data.get("qualifications", ""),
            credit_rating=data.get("credit_rating", "A"),
            price_valid_days=int(data.get("price_valid_days", 30)),
            payment_terms=data.get("payment_terms", ""),
            delivery_cycle=data.get("delivery_cycle", ""),
            after_sales=data.get("after_sales", ""),
            warranty_days=int(data.get("warranty_days", 0)),
            status=data.get("status", "pending"),
            remark=data.get("remark", ""),
        )
        supplier_id = database.create_supplier(supplier)
        return jsonify({"success": True, "supplier": database.get_supplier_by_id(supplier_id).to_dict()})

    @app.route("/api/suppliers/<int:supplier_id>", methods=["GET"])
    def api_suppliers_get(supplier_id):
        auth_check = require_auth()
        if auth_check:
            return auth_check

        supplier = database.get_supplier_by_id(supplier_id)
        if not supplier:
            return jsonify({"error": "供应商不存在"}), 404

        products = database.get_supplier_products(supplier_id=supplier_id)
        return jsonify({
            "supplier": supplier.to_dict(),
            "products": [p.to_dict() for p in products],
        })

    @app.route("/api/suppliers/<int:supplier_id>", methods=["PUT"])
    def api_suppliers_update(supplier_id):
        auth_check = require_auth("admin")
        if auth_check:
            return auth_check

        data = request.get_json() or {}
        update_fields = {}
        for field in ["name", "contact_name", "contact_phone", "contact_email", "address",
                      "business_license", "qualifications", "credit_rating", "price_valid_days",
                      "payment_terms", "delivery_cycle", "after_sales", "warranty_days", "status", "remark"]:
            if field in data:
                update_fields[field] = data[field]

        database.update_supplier(supplier_id, **update_fields)
        supplier = database.get_supplier_by_id(supplier_id)
        return jsonify({"success": True, "supplier": supplier.to_dict()})

    @app.route("/api/suppliers/<int:supplier_id>/approve", methods=["POST"])
    def api_suppliers_approve(supplier_id):
        auth_check = require_auth("admin")
        if auth_check:
            return auth_check

        database.update_supplier(supplier_id, status="approved")
        supplier = database.get_supplier_by_id(supplier_id)

        users = database.get_users(role="supplier")
        for u in users:
            if u.supplier_id == supplier_id:
                database.update_user(u.id, status="active")

        return jsonify({"success": True, "supplier": supplier.to_dict()})

    @app.route("/api/suppliers/<int:supplier_id>/reject", methods=["POST"])
    def api_suppliers_reject(supplier_id):
        auth_check = require_auth("admin")
        if auth_check:
            return auth_check

        data = request.get_json() or {}
        remark = data.get("remark", "")

        database.update_supplier(supplier_id, status="rejected", remark=remark)
        supplier = database.get_supplier_by_id(supplier_id)

        users = database.get_users(role="supplier")
        for u in users:
            if u.supplier_id == supplier_id:
                database.update_user(u.id, status="disabled")

        return jsonify({"success": True, "supplier": supplier.to_dict()})

    @app.route("/api/categories", methods=["GET"])
    def api_categories():
        from storage.categories import _normalize_type
        product_type = request.args.get("type", "goods")
        norm_type = _normalize_type(product_type)
        main_category = request.args.get("main_category")

        if main_category:
            sub_categories = get_sub_categories(main_category, product_type)
            return jsonify({
                "main_category": main_category,
                "sub_categories": sub_categories
            })

        main_categories = get_main_categories(product_type)
        return jsonify({
            "product_type": product_type,
            "main_categories": main_categories,
            "categories": CATEGORIES.get(norm_type, {})
        })

    @app.route("/api/supplier-products", methods=["GET"])
    def api_supplier_products_list():
        auth_check = require_auth()
        if auth_check:
            return auth_check

        supplier_id = request.args.get("supplier_id")
        supplier_id = int(supplier_id) if supplier_id else None
        product_type = request.args.get("product_type") or None
        main_category = request.args.get("main_category") or None
        sub_category = request.args.get("sub_category") or None
        keyword = request.args.get("keyword") or None
        limit = int(request.args.get("limit", 100))

        products = database.get_supplier_products(
            supplier_id=supplier_id,
            product_type=product_type,
            main_category=main_category,
            sub_category=sub_category,
            keyword=keyword,
            limit=limit,
        )

        suppliers = database.get_suppliers(limit=200)
        supplier_map = {s.id: s.name for s in suppliers}

        result = []
        for p in products:
            d = p.to_dict()
            d["supplier_name"] = supplier_map.get(p.supplier_id, "")
            result.append(d)

        return jsonify({"products": result})

    @app.route("/api/supplier-products", methods=["POST"])
    def api_supplier_products_create():
        auth_check = require_auth("admin")
        if auth_check:
            return auth_check

        data = request.get_json() or {}
        product = SupplierProduct(
            supplier_id=int(data.get("supplier_id", 0)),
            product_type=data.get("product_type", "goods"),
            main_category=data.get("main_category", ""),
            sub_category=data.get("sub_category", ""),
            title=data.get("title", ""),
            spec=data.get("spec", ""),
            unit=data.get("unit", ""),
            price=float(data.get("price", 0)),
            min_order=int(data.get("min_order", 1)),
            bulk_discount=data.get("bulk_discount", ""),
            delivery_cycle=data.get("delivery_cycle", ""),
            warranty_days=int(data.get("warranty_days", 0)),
            description=data.get("description", ""),
            image_url=data.get("image_url", ""),
            keyword=data.get("keyword", ""),
        )
        product_id = database.create_supplier_product(product)
        return jsonify({"success": True, "product": database.get_supplier_product_by_id(product_id).to_dict()})

    @app.route("/api/supplier-products/<int:product_id>", methods=["GET"])
    def api_supplier_products_get(product_id):
        auth_check = require_auth()
        if auth_check:
            return auth_check

        product = database.get_supplier_product_by_id(product_id)
        if not product:
            return jsonify({"error": "商品不存在"}), 404
        return jsonify({"product": product.to_dict()})

    @app.route("/api/supplier-products/<int:product_id>", methods=["PUT"])
    def api_supplier_products_update(product_id):
        auth_check = require_auth("admin")
        if auth_check:
            return auth_check

        data = request.get_json() or {}
        update_fields = {}
        for field in ["supplier_id", "product_type", "main_category", "sub_category",
                      "title", "spec", "unit", "price", "min_order",
                      "bulk_discount", "delivery_cycle", "warranty_days", "description",
                      "image_url", "status", "keyword"]:
            if field in data:
                update_fields[field] = data[field]

        database.update_supplier_product(product_id, **update_fields)
        product = database.get_supplier_product_by_id(product_id)
        return jsonify({"success": True, "product": product.to_dict()})

    @app.route("/api/supplier-products/<int:product_id>", methods=["DELETE"])
    def api_supplier_products_delete(product_id):
        auth_check = require_auth("admin")
        if auth_check:
            return auth_check

        database.update_supplier_product(product_id, status="deleted")
        return jsonify({"success": True})

    @app.route("/api/supplier-products/import/template")
    def api_supplier_products_template():
        auth_check = require_auth("admin")
        if auth_check:
            return auth_check

        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = Workbook()
        ws = wb.active
        ws.title = "供应商商品导入模板"

        headers = [
            "供应商名称*", "商品类型*", "一级分类", "二级分类",
            "商品名称*", "规格型号", "单位", "价格*", "最小起订量",
            "批量折扣", "交货周期", "质保天数", "商品描述", "图片链接",
            "关键词", "状态"
        ]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="6366f1", end_color="6366f1", fill_type="solid")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        examples = [
            ["示例供应商A", "goods", "办公设备", "打印机", "A4彩色激光打印机", "HP M254dw", "台", 2999.00, 1, "10台以上95折", "3-5工作日", 365, "高速彩色打印，支持双面", "", "打印机", "active"],
            ["示例供应商B", "service", "运维服务", "网络维护", "企业网络年度维护服务", "100节点以内", "年", 12000.00, 1, "", "7x24小时响应", 0, "包含网络设备巡检、故障排除", "", "网络维护", "active"],
        ]
        for row_idx, example in enumerate(examples, 2):
            for col_idx, value in enumerate(example, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        col_widths = [18, 10, 12, 12, 30, 20, 8, 10, 12, 18, 15, 10, 30, 30, 15, 10]
        for i, width in enumerate(col_widths, 1):
            ws.column_dimensions[chr(64 + i)].width = width

        ws.freeze_panes = "A2"

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        from flask import send_file
        return send_file(
            buf,
            as_attachment=True,
            download_name="供应商商品导入模板.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    @app.route("/api/supplier-products/import", methods=["POST"])
    def api_supplier_products_import():
        auth_check = require_auth("admin")
        if auth_check:
            return auth_check

        try:
            if "file" not in request.files:
                return jsonify({"error": "请上传文件"}), 400

            file = request.files["file"]
            if file.filename == "":
                return jsonify({"error": "请选择文件"}), 400

            if not file.filename.endswith((".xlsx", ".xls", ".csv")):
                return jsonify({"error": "仅支持 Excel (.xlsx, .xls) 和 CSV 文件"}), 400

            filename = file.filename.lower()
            products = []
            errors = []

            if filename.endswith(".csv"):
                import csv
                import io
                content = file.read().decode("utf-8-sig")
                reader = csv.DictReader(io.StringIO(content))
                for idx, row in enumerate(reader, 1):
                    product, err = _parse_supplier_product_row(row, idx)
                    if product:
                        products.append(product)
                    if err:
                        errors.append(err)
            else:
                from openpyxl import load_workbook
                wb = load_workbook(file)
                ws = wb.active

                headers = []
                for cell in ws[1]:
                    headers.append(str(cell.value).strip() if cell.value else "")

                required_headers = ["供应商名称", "商品类型", "商品名称", "价格"]
                matched = [h for h in required_headers if any(h in header for header in headers)]
                if len(matched) < 2:
                    return jsonify({
                        "error": "文件格式不正确，表头与模板不匹配。请先下载导入模板，按模板格式填写数据后再上传。",
                        "errors": [f"检测到的表头: {', '.join(headers[:10])}"]
                    }), 400

                for row_idx in range(2, ws.max_row + 1):
                    row_data = {}
                    for col_idx, header in enumerate(headers):
                        cell_value = ws.cell(row=row_idx, column=col_idx + 1).value
                        row_data[header] = str(cell_value).strip() if cell_value is not None else ""

                    if not any(row_data.values()):
                        continue

                    product, err = _parse_supplier_product_row(row_data, row_idx)
                    if product:
                        products.append(product)
                    if err:
                        errors.append(err)

            if not products:
                if errors:
                    return jsonify({
                        "error": "没有解析到有效的商品数据",
                        "errors": errors
                    }), 400
                return jsonify({
                    "error": "文件中没有数据行，请填写数据后再上传。"
                }), 400

            success_count, db_errors = database.batch_insert_supplier_products(products)
            errors.extend(db_errors)

            return jsonify({
                "success": True,
                "total": len(products) + len([e for e in errors if "解析失败" in e or "缺少" in e]),
                "success_count": success_count,
                "error_count": len(errors),
                "errors": errors[:50]
            })

        except Exception as e:
            return jsonify({"error": f"导入失败: {str(e)}"}), 500

    def _parse_supplier_product_row(row, row_num):
        from storage.models import SupplierProduct

        supplier_name = row.get("供应商名称*", "") or row.get("供应商名称", "")
        product_type = row.get("商品类型*", "") or row.get("商品类型", "")
        title = row.get("商品名称*", "") or row.get("商品名称", "")
        price_str = row.get("价格*", "") or row.get("价格", "")

        if not supplier_name:
            return None, f"第{row_num}行: 缺少供应商名称"
        if not title:
            return None, f"第{row_num}行: 缺少商品名称"
        if not price_str:
            return None, f"第{row_num}行: 缺少价格"

        supplier = database.get_supplier_by_name(supplier_name)
        if not supplier:
            return None, f"第{row_num}行: 供应商 '{supplier_name}' 不存在"

        try:
            price = float(price_str)
        except (ValueError, TypeError):
            return None, f"第{row_num}行: 价格格式不正确 '{price_str}'"

        product_type = product_type if product_type in ("goods", "service") else "goods"

        try:
            min_order = int(row.get("最小起订量", "1") or 1)
        except (ValueError, TypeError):
            min_order = 1

        try:
            warranty_days = int(row.get("质保天数", "0") or 0)
        except (ValueError, TypeError):
            warranty_days = 0

        status = row.get("状态", "active") or "active"
        if status not in ("active", "inactive"):
            status = "active"

        product = SupplierProduct(
            supplier_id=supplier.id,
            product_type=product_type,
            main_category=row.get("一级分类", "") or "",
            sub_category=row.get("二级分类", "") or "",
            title=title,
            spec=row.get("规格型号", "") or "",
            unit=row.get("单位", "") or "",
            price=price,
            min_order=min_order,
            bulk_discount=row.get("批量折扣", "") or "",
            delivery_cycle=row.get("交货周期", "") or "",
            warranty_days=warranty_days,
            description=row.get("商品描述", "") or "",
            image_url=row.get("图片链接", "") or "",
            status=status,
            keyword=row.get("关键词", "") or "",
        )

        return product, None

    @app.route("/api/rfq", methods=["GET"])
    def api_rfq_list():
        auth_check = require_auth()
        if auth_check:
            return auth_check

        user = get_current_user()
        status = request.args.get("status")
        limit = int(request.args.get("limit", 50))

        if user.role == "admin":
            rfqs = database.get_rfq_records(status=status, limit=limit)
        else:
            rfqs = database.get_rfq_records(user_id=user.id, status=status, limit=limit)

        return jsonify({"rfqs": [r.to_dict() for r in rfqs]})

    @app.route("/api/rfq", methods=["POST"])
    def api_rfq_create():
        auth_check = require_auth()
        if auth_check:
            return auth_check

        user = get_current_user()
        data = request.get_json() or {}
        rfq = RFQRecord(
            user_id=user.id,
            product_type=data.get("product_type", "goods"),
            title=data.get("title", ""),
            spec=data.get("spec", ""),
            quantity=int(data.get("quantity", 1)),
            unit=data.get("unit", ""),
            expected_price=float(data.get("expected_price", 0)),
            delivery_requirement=data.get("delivery_requirement", ""),
            status="pending",
        )
        rfq_id = database.create_rfq(rfq)
        return jsonify({"success": True, "rfq": database.get_rfq_by_id(rfq_id).to_dict()})

    @app.route("/api/rfq/<int:rfq_id>", methods=["GET"])
    def api_rfq_get(rfq_id):
        auth_check = require_auth()
        if auth_check:
            return auth_check

        rfq = database.get_rfq_by_id(rfq_id)
        if not rfq:
            return jsonify({"error": "询价单不存在"}), 404

        user = get_current_user()
        if user.role != "admin" and rfq.user_id != user.id:
            return jsonify({"error": "无权查看此询价单"}), 403

        return jsonify({"rfq": rfq.to_dict()})

    @app.route("/api/rfq/<int:rfq_id>", methods=["PUT"])
    def api_rfq_update(rfq_id):
        auth_check = require_auth()
        if auth_check:
            return auth_check

        rfq = database.get_rfq_by_id(rfq_id)
        if not rfq:
            return jsonify({"error": "询价单不存在"}), 404

        user = get_current_user()
        if user.role != "admin" and rfq.user_id != user.id:
            return jsonify({"error": "无权修改此询价单"}), 403

        data = request.get_json() or {}
        update_fields = {}
        for field in ["product_type", "title", "spec", "quantity", "unit", "expected_price",
                      "delivery_requirement", "status", "assigned_supplier_id",
                      "supplier_quote", "quote_response"]:
            if field in data:
                update_fields[field] = data[field]

        database.update_rfq(rfq_id, **update_fields)
        rfq = database.get_rfq_by_id(rfq_id)
        return jsonify({"success": True, "rfq": rfq.to_dict()})

    @app.route("/api/rfq/<int:rfq_id>/quote", methods=["POST"])
    def api_rfq_quote(rfq_id):
        auth_check = require_auth("admin")
        if auth_check:
            return auth_check

        data = request.get_json() or {}
        database.update_rfq(
            rfq_id,
            status="quoted",
            supplier_quote=float(data.get("supplier_quote", 0)),
            quote_response=data.get("quote_response", ""),
        )
        rfq = database.get_rfq_by_id(rfq_id)
        return jsonify({"success": True, "rfq": rfq.to_dict()})

    @app.route("/api/deploy", methods=["POST"])
    def api_deploy():
        import subprocess
        try:
            pull = subprocess.run(
                ["git", "pull", "origin", "trae/agent-w4LV0u"],
                capture_output=True, text=True, timeout=30,
                cwd="/www/bijia-system"
            )
            install = subprocess.run(
                ["bash", "-c", "source /www/bijia-system/price-compare/venv/bin/activate && pip install -r /www/bijia-system/price-compare/requirements.txt"],
                capture_output=True, text=True, timeout=120
            )
            restart = subprocess.run(
                ["systemctl", "restart", "bijia"],
                capture_output=True, text=True, timeout=10
            )
            return jsonify({
                "success": True,
                "pull_output": pull.stdout,
                "pull_errors": pull.stderr,
                "install_output": install.stdout[-500:] if install.stdout else "",
                "install_errors": install.stderr[-500:] if install.stderr else "",
                "restart_output": restart.stdout,
            })
        except Exception as e:
            return jsonify({"success": False, "error": str(e)}), 500

    return app


if __name__ == "__main__":
    app = create_app()
    app.run(host="0.0.0.0", port=5000, debug=True)